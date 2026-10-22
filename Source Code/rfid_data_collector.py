"""
Program Pengambilan Data RSSI UHF-RFID
Reader     : EL-UHF-RC4-2 (Electron Indonesia)
Protokol   : Binary frame, HEAD=0xCF, 115200 8N1

Struktur program
1. CONFIGURATION
2. RFID PROTOCOL
3. RFID READER
4. DISPLAY / UI
5. DATA COLLECTION
6. DATA PROCESSING
7. EXCEL EXPORT
8. MAIN
"""

# ============================================================#
#                        CONFIGURATION                        #
# ============================================================#

import time
from datetime import datetime
from typing import Dict, List, Optional, Tuple

import serial
from openpyxl import Workbook

PORT = "COM11"
BAUDRATE = 115200
TARGET_SAMPLES = 100

# EPC tag referensi.
TAG_REFERENSI = {
    "R1": {
        "epc": "E2 80 69 95 00 00 40 17 60 38 40 F7",
    },
    "R2": {
        "epc": "E2 80 69 95 00 00 50 17 60 38 44 F7",
    },
    "R3": {
        "epc": "E2 80 69 95 00 00 50 17 60 37 90 F7",
    },
}

# Konfigurasi posisi tag referensi per skenario.
SCENARIO_CONFIG = {
    "Skenario1": {
        "R1": 50,
        "R2": 200,
    },
    "Skenario2": {
        "R1": 50,
        "R2": 125,
        "R3": 200,
    },
}

TAG_TARGET_EPC = "E2 80 69 95 00 00 40 17 60 38 34 F7"

CONDITION_MAP = {
    "1": "LOS",
    "2": "NLOS_Satu_Halangan",
    "3": "NLOS_Tiga_Halangan",
    "4": "NLOS_Lima_Halangan",
}

# ============================================================#
#                        RFID PROTOCOL                        #
# ============================================================#

HEADER = 0xCF
READER_ADDR = 0x00
FRAME_INVENTORY_START = bytes.fromhex("CFFF0001050000000000F5B5")
CMD_INVENTORY_STOP = bytearray([0x00, 0x02])
CMD_INIT = bytearray([0x00, 0x50])

def calculate_crc16(data: bytes) -> bytearray:
    """Hitung CRC16 sesuai protokol reader."""
    value = 0xFFFF
    for byte in data:
        value ^= byte
        for _ in range(8):
            value = (value >> 1) ^ 0x8408 if value & 0x0001 else value >> 1
    return bytearray([value >> 8, value & 0xFF])

def build_command(cmd: bytearray, param: Optional[bytearray] = None) -> bytes:
    """Bangun frame command untuk dikirim ke reader."""
    if param is None:
        param = bytearray()

    frame = bytearray([HEADER, READER_ADDR])
    frame.extend(cmd)
    frame.append(len(param))
    frame.extend(param)
    frame.extend(calculate_crc16(frame))
    return bytes(frame)

def parse_rssi(byte_high: int, byte_low: int) -> float:
    """Parse RSSI 2-byte signed integer dengan resolusi 0.1 dBm."""
    raw = (byte_high << 8) | byte_low
    signed = raw - 0x10000 if raw >= 0x8000 else raw
    return signed / 10.0

def parse_inventory_response(data: bytes) -> List[dict]:
    """Parse seluruh frame inventory yang terdapat pada buffer serial."""
    tags = []
    index = 0

    while index < len(data):
        if data[index] != HEADER:
            index += 1
            continue

        if index + 6 >= len(data):
            break

        cmd_word = (data[index + 2] << 8) | data[index + 3]
        length = data[index + 4]
        total_length = 5 + length + 2

        if index + total_length > len(data):
            break

        status = data[index + 5]
        param = data[index + 6:index + 5 + length]

        if cmd_word == 0x0001 and status == 0x00 and len(param) >= 5:
            rssi_dbm = parse_rssi(param[0], param[1])
            tag_info_len = param[4]
            tag_info = param[5:5 + tag_info_len]
            epc_hex = " ".join(f"{byte:02X}" for byte in tag_info)

            if epc_hex:
                tags.append({
                    "epc": epc_hex,
                    "rssi_dbm": rssi_dbm,
                })

        index += total_length

    return tags

# ============================================================#
#                         RFID READER                         #
# ============================================================#

class RFIDReader:
    """Wrapper komunikasi serial dengan reader EL-UHF-RC4-2."""

    def __init__(self, port: str, baudrate: int = BAUDRATE):
        self.ser: Optional[serial.Serial] = None
        self.port = port
        self.baudrate = baudrate

    def connect(self) -> bool:
        """Buka koneksi serial."""
        try:
            self.ser = serial.Serial(
                port=self.port,
                baudrate=self.baudrate,
                bytesize=serial.EIGHTBITS,
                parity=serial.PARITY_NONE,
                stopbits=serial.STOPBITS_ONE,
                timeout=1.0,
            )
            time.sleep(0.5)
            self.ser.reset_input_buffer()
            return True
        except serial.SerialException as exc:
            print_error(f"Gagal terhubung ke {self.port}: {exc}")
            return False

    def init(self) -> None:
        """Inisialisasi reader."""
        if self.ser is None:
            raise RuntimeError("Reader belum terhubung.")
        self.ser.write(build_command(CMD_INIT))
        time.sleep(0.3)
        self.ser.reset_input_buffer()

    def start(self) -> None:
        """Mulai inventory/tag scanning."""
        if self.ser is None:
            raise RuntimeError("Reader belum terhubung.")
        self.ser.write(FRAME_INVENTORY_START)
        time.sleep(0.05)

    def stop(self) -> None:
        """Hentikan inventory/tag scanning."""
        if self.ser is None:
            return
        try:
            self.ser.write(build_command(CMD_INVENTORY_STOP))
            time.sleep(0.2)
            self.ser.reset_input_buffer()
        except (serial.SerialException, AttributeError):
            pass

    def read_buffer(self, timeout: float = 0.2) -> bytes:
        """Baca seluruh data yang masuk selama window timeout."""
        if self.ser is None:
            return b""

        buffer = bytearray()
        start_time = time.time()

        while time.time() - start_time < timeout:
            waiting = self.ser.in_waiting
            if waiting:
                buffer.extend(self.ser.read(waiting))
            else:
                time.sleep(0.005)

        return bytes(buffer)

    def disconnect(self) -> None:
        """Hentikan scanning dan tutup koneksi serial."""
        if self.ser is None:
            return
        try:
            self.stop()
            if self.ser.is_open:
                self.ser.close()
        except (serial.SerialException, AttributeError):
            pass
        finally:
            self.ser = None

# ============================================================#
#                        DISPLAY / UI                         #
# ============================================================#

def print_header(title: str) -> None:
    line = "═" * 62
    print(f"\n╔{line}╗")
    print(f"║            {title:<50}║")
    print(f"╚{line}╝")

def print_separator() -> None:
    print("  " + "─" * 60)

def print_error(message: str) -> None:
    print(f"\n  [ERROR] {message}")

def print_config(
    skenario: str,
    kondisi: str,
    active_tags: dict,
    jarak_cm: Optional[float] = None,
) -> None:
    """Tampilkan konfigurasi pengambilan data."""
    print(f"  Skenario      : {skenario}")
    print(f"  Kondisi       : {kondisi}")

    if jarak_cm is not None:
        print(f"  Jarak Target  : {jarak_cm:g} cm")

    print("  Tag Referensi :")
    for name, config in active_tags.items():
        print(f"    {name:<3} : {config['jarak_cm']} cm")

    print(f"  Sampel Target : {TARGET_SAMPLES}")

def print_progress(count: int, total: int, width: int = 24) -> str:
    """Buat progress bar pengambilan, bukan bar kualitas sinyal."""
    count = min(count, total)
    percent = (count / total * 100) if total else 0.0
    filled = int(percent / 100 * width)
    bar = "█" * filled + "░" * (width - filled)
    return f"[{bar}] {count}/{total} ({percent:.1f}%)"

def print_statistics_line(name: str, records: List[dict]) -> None:
    """Tampilkan avg/min/max RSSI sebuah tag."""
    stats = calculate_statistics(records)
    if stats["count"] == 0:
        print(f"  {name:<6}: belum ada bacaan")
        return

    print(
        f"  RSSI {name:<6}: avg={stats['avg']:.2f} | "
        f"min={stats['min']:.1f} | max={stats['max']:.1f} dBm"
    )

def print_progress_status(
    records: dict,
    tag_config: dict,
    jarak_cm: float,
) -> None:
    """Tampilkan status target dan statistik seluruh tag secara berkala."""
    print("\n  ┌─ STATUS PENGAMBILAN ────────────────────────────────────┐")
    progress = print_progress(len(records["target"]), TARGET_SAMPLES, 20)
    print(f"  │ Target : {progress:<42}     │")

    print("  │                                                         │")
    for name in tag_config:
        count = len(records.get(name, []))
        print(f"  │ {name:<6}: {count:>4} bacaan{'':<37}│")

    print("  │                                                         │")

    def format_stats(records_for_tag: List[dict]) -> str:
        stats = calculate_statistics(records_for_tag)
        if stats["count"] == 0:
            return "belum ada bacaan"
        return (
            f"avg={stats['avg']:.2f} | "
            f"min={stats['min']:.1f} | max={stats['max']:.1f} dBm"
        )

    print(f"  │ RSSI Target : {format_stats(records['target']):<42}│")
    for name in tag_config:
        print(f"  │ RSSI {name:<6}: {format_stats(records[name]):<42} │")

    print("  └─────────────────────────────────────────────────────────┘")

# ============================================================#
#                      DATA COLLECTION                        #
# ============================================================#

def run_diagnostic(reader: RFIDReader) -> None:
    """Scan semua EPC yang terbaca sampai pengguna menekan Ctrl+C."""
    print_header("MODE DIAGNOSTIK")
    print("  Scan semua tag yang terbaca oleh reader.")
    print("  Tekan Ctrl+C untuk berhenti.\n")

    seen: Dict[str, List[float]] = {}
    reader.start()

    try:
        while True:
            raw = reader.read_buffer(0.5)
            tags = parse_inventory_response(raw)

            for tag in tags:
                epc = tag["epc"]
                rssi = tag["rssi_dbm"]
                timestamp = datetime.now().strftime("%H:%M:%S.%f")[:-3]

                if epc not in seen:
                    seen[epc] = []
                    print("  + TAG BARU")
                    print(f"    EPC : {epc}")

                seen[epc].append(rssi)
                print(
                    f"  [{timestamp}] EPC ...{epc[-11:]} | "
                    f"RSSI {rssi:>6.1f} dBm | n={len(seen[epc])}"
                )

    except KeyboardInterrupt:
        print("\n")
        print_separator()
        print("  RINGKASAN TAG TERDETEKSI")
        print_separator()

        for epc, values in seen.items():
            print(f"  EPC : {epc}")
            print(f"  N   : {len(values)} bacaan")
            print(f"  Avg : {sum(values) / len(values):.2f} dBm")
            print(f"  Min : {min(values):.1f} dBm")
            print(f"  Max : {max(values):.1f} dBm")
            print_separator()

        print("  Salin EPC ke konfigurasi tag referensi / tag target.")

    finally:
        reader.stop()

def collect_reference(
    reader: RFIDReader,
    tag_config: dict,
    kondisi: str,
    skenario: str,
) -> dict:
    """Ambil data seluruh tag referensi secara bersamaan."""
    records = {name: [] for name in tag_config}
    epc_to_name = {config["epc"]: name for name, config in tag_config.items()}

    print_header(f"PENGAMBILAN DATA REFERENSI — {skenario} | {kondisi}")
    print_config(skenario, kondisi, active_tags=tag_config)
    print("  Semua tag referensi harus sudah berada di posisi masing-masing.")
    input("  Tekan ENTER untuk mulai...\n")

    print_separator()
    print(f"  {'Waktu':<15} {'Tag':<6} {'Jarak':>7}  {'RSSI':>8}  {'Progress'}")
    print_separator()

    reader.start()

    try:
        while not all(len(records[name]) >= TARGET_SAMPLES for name in tag_config):
            raw = reader.read_buffer(0.2)
            tags = parse_inventory_response(raw)

            for tag in tags:
                epc = tag["epc"]
                rssi = tag["rssi_dbm"]
                name = epc_to_name.get(epc)
                timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]

                if name is None:
                    print(f"  [?] UNKNOWN | EPC {epc} | RSSI {rssi:.1f} dBm")
                    continue

                if len(records[name]) >= TARGET_SAMPLES:
                    continue

                records[name].append({
                    "epc": epc,
                    "rssi_dbm": rssi,
                    "timestamp": timestamp,
                })

                count = len(records[name])
                distance = tag_config[name]["jarak_cm"]

                print(
                    f"  [{datetime.now().strftime('%H:%M:%S.%f')[:-3]}] "
                    f"{name:<6} {distance:>5g}cm  {rssi:>7.1f}  "
                    f"{count}/{TARGET_SAMPLES}"
                )

    except KeyboardInterrupt:
        print("\n  [INFO] Pengambilan dihentikan pengguna.")
    finally:
        reader.stop()

    print_separator()
    print("  RINGKASAN HASIL")
    print_separator()
    for name in tag_config:
        print_statistics_line(name, records[name])
    print_separator()

    return records

def collect_target(
    reader: RFIDReader,
    epc_target: str,
    jarak_cm: float,
    kondisi: str,
    skenario: str,
    tag_config: dict,
) -> dict:
    """
    Ambil data tag target dan seluruh tag referensi dalam satu sesi.

    Target berhenti pada TARGET_SAMPLES. Pembacaan referensi disimpan selama
    sesi berlangsung, sehingga jumlah bacaan referensi dapat berbeda dari
    jumlah sampel target, seperti dataset penelitian sebelumnya.
    """
    records = {"target": []}
    for name in tag_config:
        records[name] = []

    epc_to_name = {config["epc"]: name for name, config in tag_config.items()}

    print_header(f"PENGAMBILAN DATA TARGET — {skenario} | {kondisi}")
    print_config(
        skenario,
        kondisi,
        active_tags=tag_config,
        jarak_cm=jarak_cm,
    )
    print("  Tag referensi tetap di posisinya dan tidak dipindahkan.")
    input("  Tekan ENTER untuk mulai...\n")

    print_separator()
    print(f"  {'Waktu':<15} {'Tag':<8} {'Jarak':>7}  {'RSSI':>8}  {'Status'}")
    print_separator()

    reader.start()
    last_status_time = time.time()

    try:
        while len(records["target"]) < TARGET_SAMPLES:
            raw = reader.read_buffer(0.2)
            tags = parse_inventory_response(raw)

            for tag in tags:
                epc = tag["epc"]
                rssi = tag["rssi_dbm"]
                now_full = datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
                now_short = datetime.now().strftime("%H:%M:%S.%f")[:-3]

                if epc == epc_target:
                    if len(records["target"]) < TARGET_SAMPLES:
                        records["target"].append({
                            "rssi_dbm": rssi,
                            "timestamp": now_full,
                        })

                    print(
                        f"  [{now_short}] TARGET  {jarak_cm:>5g}cm  "
                        f"{rssi:>7.1f}  "
                        f"{len(records['target'])}/{TARGET_SAMPLES}"
                    )

                elif epc in epc_to_name:
                    ref_name = epc_to_name[epc]
                    records[ref_name].append({
                        "rssi_dbm": rssi,
                        "timestamp": now_full,
                    })

                else:
                    print(f"  [?] UNKNOWN | EPC {epc} | RSSI {rssi:.1f} dBm")

                if len(records["target"]) >= TARGET_SAMPLES:
                    break

            if time.time() - last_status_time >= 2.0:
                print_progress_status(records, tag_config, jarak_cm)
                last_status_time = time.time()

    except KeyboardInterrupt:
        print("\n  [INFO] Pengambilan dihentikan pengguna.")
    finally:
        reader.stop()

    print_separator()
    print("  DATA SELESAI")
    print_separator()
    print(f"  Target : {len(records['target'])}/{TARGET_SAMPLES} sampel")
    for ref_name in tag_config:
        print(f"  {ref_name:<6}: {len(records[ref_name])} bacaan referensi")

    print_statistics_line("Target", records["target"])
    for ref_name in tag_config:
        print_statistics_line(ref_name, records[ref_name])
    print_separator()

    return records

# ============================================================#
#                      DATA PROCESSING                        #
# ============================================================#

def calculate_statistics(records: List[dict]) -> dict:
    """Hitung jumlah, rata-rata, minimum, dan maksimum RSSI."""
    if not records:
        return {"count": 0, "avg": 0.0, "min": 0.0, "max": 0.0}

    values = [item["rssi_dbm"] for item in records]
    return {
        "count": len(values),
        "avg": sum(values) / len(values),
        "min": min(values),
        "max": max(values),
    }

def build_snapshot(
    target: Optional[dict] = None,
    references: Optional[dict] = None,
) -> dict:
    """Bangun snapshot internal tanpa mengubah format Excel lama."""
    return {
        "timestamp": target.get("timestamp") if target else None,
        "target": target,
        "references": references or {},
    }

# ============================================================#
#                         EXCEL EXPORT                        #
# ============================================================#

def set_column_widths(ws) -> None:
    """Atur lebar kolom agar Excel mudah dibaca."""
    for column in ws.columns:
        width = max((len(str(cell.value or "")) for cell in column), default=0)
        ws.column_dimensions[column[0].column_letter].width = min(width + 4, 40)

def save_reference(
    records: dict,
    tag_config: dict,
    kondisi: str,
    skenario: str,
) -> str:
    """Simpan data referensi dengan struktur Excel dataset lama."""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"data_{skenario}_{kondisi}_Referensi_{timestamp}.xlsx"

    workbook = Workbook()
    workbook.remove(workbook.active)

    for name, recs in records.items():
        if not recs:
            continue

        sheet = workbook.create_sheet(title=name)
        distance = tag_config[name]["jarak_cm"]

        sheet.append([
            "No",
            "Timestamp",
            "EPC",
            "RSSI (dBm)",
            "Jarak Aktual (cm)",
            "Label Tag",
            "Kondisi",
            "Skenario",
        ])

        for index, record in enumerate(recs, start=1):
            sheet.append([
                index,
                record["timestamp"],
                record["epc"],
                record["rssi_dbm"],
                distance,
                name,
                kondisi,
                skenario,
            ])

        set_column_widths(sheet)

    summary = workbook.create_sheet(title="Ringkasan", index=0)
    summary.append([
        "Tag",
        "Jarak (cm)",
        "Jumlah Sampel",
        "RSSI Avg",
        "RSSI Min",
        "RSSI Max",
        "Kondisi",
        "Skenario",
    ])

    for name, recs in records.items():
        if not recs:
            continue

        stats = calculate_statistics(recs)
        summary.append([
            name,
            tag_config[name]["jarak_cm"],
            stats["count"],
            round(stats["avg"], 2),
            stats["min"],
            stats["max"],
            kondisi,
            skenario,
        ])

    set_column_widths(summary)
    workbook.save(filename)

    print(f"\n  [OK] File tersimpan: {filename}")
    print(f"  Sheet: Ringkasan + {' + '.join(records.keys())}")
    return filename

def save_target(
    records: dict,
    jarak_cm: float,
    kondisi: str,
    skenario: str,
    tag_config: dict,
) -> str:
    """Simpan target dan referensi dengan struktur Excel dataset lama."""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"data_{skenario}_{kondisi}_Target_{int(jarak_cm)}cm_{timestamp}.xlsx"

    workbook = Workbook()
    workbook.remove(workbook.active)

    all_tags = {"target": ("Target", jarak_cm)}
    for name, config in tag_config.items():
        all_tags[name] = (name, config["jarak_cm"])

    for key, (label, distance) in all_tags.items():
        recs = records.get(key, [])
        if not recs:
            continue

        sheet = workbook.create_sheet(title=label)
        sheet.append([
            "No",
            "Timestamp",
            "RSSI (dBm)",
            "Jarak Aktual (cm)",
            "Label Tag",
            "Kondisi",
            "Skenario",
        ])

        for index, record in enumerate(recs, start=1):
            sheet.append([
                index,
                record.get("timestamp", ""),
                record["rssi_dbm"],
                distance,
                label,
                kondisi,
                skenario,
            ])

        set_column_widths(sheet)

    summary = workbook.create_sheet(title="Ringkasan", index=0)
    summary.append([
        "Tag",
        "Jarak (cm)",
        "Jumlah Sampel",
        "RSSI Avg",
        "RSSI Min",
        "RSSI Max",
        "Kondisi",
        "Skenario",
    ])

    for key, (label, distance) in all_tags.items():
        recs = records.get(key, [])
        if not recs:
            continue

        stats = calculate_statistics(recs)
        summary.append([
            label,
            distance,
            stats["count"],
            round(stats["avg"], 2),
            stats["min"],
            stats["max"],
            kondisi,
            skenario,
        ])

    set_column_widths(summary)
    workbook.save(filename)

    print(f"\n  [OK] File tersimpan: {filename}")
    print(f"  Sheet: Ringkasan + Target + {' + '.join(tag_config.keys())}")
    return filename

# ============================================================#
#                              MAIN                           #
# ============================================================#

def show_menu() -> str:
    print_header("PROGRAM PENGAMBILAN DATA RSSI UHF-RFID")
    print("  1 = Pengambilan data TAG REFERENSI")
    print("  2 = Pengambilan data TAG TARGET")
    print("  3 = Diagnostik EPC & RSSI")
    return input("\n  Pilihan (1/2/3): ").strip()

def select_scenario() -> Tuple[str, dict]:
    print("\n  Skenario:")
    print("  1 = 2 tag referensi (R1=50 cm, R2=200 cm)")
    print("  2 = 3 tag referensi (R1=50 cm, R2=125 cm, R3=200 cm)")

    choice = input("  Pilih skenario (1/2): ").strip()
    if choice not in {"1", "2"}:
        raise ValueError("Pilihan skenario harus 1 atau 2.")

    skenario = f"Skenario{choice}"
    active_tags = {}

    for name, distance in SCENARIO_CONFIG[skenario].items():
        active_tags[name] = {
            "epc": TAG_REFERENSI[name]["epc"],
            "jarak_cm": distance,
        }

    return skenario, active_tags

def select_condition() -> str:
    print("\n  Kondisi:")
    print("  1 = LOS")
    print("  2 = NLOS Satu Halangan")
    print("  3 = NLOS Tiga Halangan")
    print("  4 = NLOS Lima Halangan")

    choice = input("  Pilih kondisi (1-4): ").strip()
    if choice not in CONDITION_MAP:
        raise ValueError("Pilihan kondisi harus 1 sampai 4.")

    return CONDITION_MAP[choice]

def main() -> None:
    mode = show_menu()

    if mode not in {"1", "2", "3"}:
        print_error("Pilihan mode tidak valid.")
        return

    reader = RFIDReader(PORT, BAUDRATE)

    if not reader.connect():
        return

    print(f"  [OK] Terhubung ke {PORT}")

    try:
        reader.init()
        print("  [OK] Reader diinisialisasi.")

        if mode == "3":
            run_diagnostic(reader)
            return

        skenario, active_tags = select_scenario()
        kondisi = select_condition()

        if mode == "1":
            records = {}
            try:
                records = collect_reference(
                    reader,
                    active_tags,
                    kondisi,
                    skenario,
                )
                save_reference(
                    records,
                    active_tags,
                    kondisi,
                    skenario,
                )
                print("  [SELESAI] Data referensi tersimpan.")
            except Exception as exc:
                print_error(str(exc))
                if records:
                    save_reference(
                        records,
                        active_tags,
                        kondisi,
                        skenario,
                    )

        elif mode == "2":
            try:
                jarak_cm = float(
                    input("\n  Jarak tag target dari reader (cm): ").strip()
                )
                if jarak_cm <= 0:
                    raise ValueError("Jarak target harus lebih besar dari 0.")

                records = collect_target(
                    reader,
                    TAG_TARGET_EPC,
                    jarak_cm,
                    kondisi,
                    skenario,
                    active_tags,
                )

                if records["target"]:
                    save_target(
                        records,
                        jarak_cm,
                        kondisi,
                        skenario,
                        active_tags,
                    )
                    print("  [SELESAI] Jalankan kembali untuk posisi target berikutnya.")
                else:
                    print("  [WARN] Tidak ada data target yang terkumpul.")

            except Exception as exc:
                print_error(str(exc))

    except KeyboardInterrupt:
        print("\n  [INFO] Program dihentikan pengguna.")
    except Exception as exc:
        print_error(str(exc))
    finally:
        reader.disconnect()
        print("  [OK] Koneksi reader ditutup.")

if __name__ == "__main__":
    main()